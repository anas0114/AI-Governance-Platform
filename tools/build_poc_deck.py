from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parents[1] / "AI-Governance-Platform-POC.pptx"

NAVY = RGBColor(16, 39, 42)
TEAL = RGBColor(8, 127, 120)
MINT = RGBColor(213, 243, 237)
INK = RGBColor(20, 33, 38)
MUTED = RGBColor(96, 113, 119)
PALE = RGBColor(243, 246, 248)
WHITE = RGBColor(255, 255, 255)
ORANGE = RGBColor(224, 130, 54)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def bg(slide, color=PALE):
    shape = slide.background
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def text(slide, value, x, y, w, h, size=18, color=INK, bold=False, font="Aptos", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def pill(slide, value, x, y, w, color=TEAL):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.34))
    shape.fill.solid(); shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    tf = shape.text_frame; tf.clear(); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = value; r.font.name = "Aptos"; r.font.size = Pt(10); r.font.bold = True; r.font.color.rgb = WHITE
    return shape


def card(slide, title, body, x, y, w, h, accent=TEAL):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = RGBColor(220, 229, 231)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
    text(slide, title, x + 0.26, y + 0.2, w - 0.45, 0.35, 16, INK, True)
    text(slide, body, x + 0.26, y + 0.68, w - 0.48, h - 0.8, 12, MUTED)


def title(slide, kicker, heading, subtitle=""):
    text(slide, kicker.upper(), 0.65, 0.45, 5, 0.25, 10, TEAL, True)
    text(slide, heading, 0.65, 0.82, 11.6, 0.65, 30, INK, True)
    if subtitle:
        text(slide, subtitle, 0.68, 1.55, 11.2, 0.45, 14, MUTED)


def footer(slide, number):
    text(slide, "AI GOVERNANCE & COMPLIANCE PLATFORM", 0.65, 7.12, 6, 0.18, 8, MUTED, True)
    text(slide, f"{number:02d}", 12.2, 7.08, 0.45, 0.22, 9, TEAL, True, align=PP_ALIGN.RIGHT)


# 1 cover
slide = prs.slides.add_slide(prs.slide_layouts[6]); bg(slide, NAVY)
shape = slide.shapes.add_shape(MSO_SHAPE.ARC, Inches(8.9), Inches(-1.2), Inches(6), Inches(6))
shape.line.color.rgb = TEAL; shape.line.width = Pt(2.5); shape.fill.background()
shape2 = slide.shapes.add_shape(MSO_SHAPE.ARC, Inches(9.6), Inches(-0.55), Inches(4.8), Inches(4.8))
shape2.line.color.rgb = ORANGE; shape2.line.width = Pt(1.5); shape2.fill.background()
pill(slide, "PROOF OF CONCEPT / 2026", 0.75, 0.78, 2.15, TEAL)
text(slide, "AI Governance\n& Compliance", 0.75, 1.65, 8.5, 1.6, 44, WHITE, True)
text(slide, "A working platform for policy, regulation, risk, compliance,\nand explainable AI-assisted governance decisions.", 0.8, 3.65, 6.8, 0.8, 18, MINT)
text(slide, "Next.js  •  FastAPI  •  PostgreSQL  •  LangChain  •  LangGraph  •  ChromaDB", 0.8, 6.45, 8.8, 0.3, 11, RGBColor(180, 213, 209))
text(slide, "01", 11.9, 6.38, 0.7, 0.5, 18, TEAL, True, align=PP_ALIGN.RIGHT)

# 2 executive snapshot
slide = prs.slides.add_slide(prs.slide_layouts[6]); bg(slide); title(slide, "The proposition", "A governance operating system for AI", "From evidence and regulations to a traceable recommendation.")
card(slide, "Govern", "Policies, frameworks, controls, approvals, and audit findings in one workspace.", 0.7, 2.35, 3.75, 2.0, TEAL)
card(slide, "Understand", "RAG over regulations, standards, findings, and governance documents with citations.", 4.8, 2.35, 3.75, 2.0, ORANGE)
card(slide, "Decide", "Specialist agents collaborate in sequence and preserve the reasoning trail.", 8.9, 2.35, 3.75, 2.0, RGBColor(91, 105, 190))
pill(slide, "WORKING FULL-STACK POC", 0.7, 5.3, 2.15, NAVY)
text(slide, "The platform has a usable dashboard, authenticated workflows, document ingestion, agent tooling, MCP interfaces, and an auditable collaboration flow.", 0.7, 5.83, 11.2, 0.6, 17, INK, True)
footer(slide, 2)

# 3 architecture
slide = prs.slides.add_slide(prs.slide_layouts[6]); bg(slide); title(slide, "Architecture", "One platform, four cooperating layers", "The same domain data powers the UI, REST API, agents, and MCP clients.")
labels = [("Next.js + Tailwind", 0.8, TEAL), ("FastAPI + SQLAlchemy", 3.45, ORANGE), ("LangGraph + LangChain", 6.1, RGBColor(91, 105, 190)), ("PostgreSQL + ChromaDB", 8.75, NAVY)]
for label, x, color in labels:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.6), Inches(2.25), Inches(1.05))
    shape.fill.solid(); shape.fill.fore_color.rgb = color; shape.line.fill.background()
    tf = shape.text_frame; tf.clear(); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label; r.font.name = "Aptos"; r.font.size = Pt(15); r.font.bold = True; r.font.color.rgb = WHITE
    if x < 8.75:
        text(slide, "→", x + 2.3, 2.88, 0.35, 0.3, 22, TEAL, True, align=PP_ALIGN.CENTER)
card(slide, "Human-facing", "Dashboard, repository, research, agents, collaboration, and document workflows.", 0.8, 4.55, 3.55, 1.35)
card(slide, "Intelligence", "RAG retrieval, citation answers, specialist tools, memory, orchestration, and observability.", 4.85, 4.55, 3.55, 1.35, ORANGE)
card(slide, "Interoperability", "Four MCP servers expose policies, compliance, audits, and documents to external clients.", 8.9, 4.55, 3.55, 1.35, RGBColor(91, 105, 190))
footer(slide, 3)

# 4 phase map
slide = prs.slides.add_slide(prs.slide_layouts[6]); bg(slide); title(slide, "Delivery map", "Five phases, one traceable flow", "Each phase adds a new decision surface without breaking the previous one.")
phases = [
    ("01", "Foundation", "Repository • Dashboard • Policies • Research", TEAL),
    ("02", "Knowledge", "Ingest • Chunk • Embed • Retrieve • Cite", ORANGE),
    ("03", "Agents", "Legal • Compliance • Risk • Advisor", RGBColor(91, 105, 190)),
    ("04", "MCP", "Policy • Compliance • Audit • Documents", NAVY),
    ("05", "Collaboration", "Research → Compliance → Risk → Governance", RGBColor(37, 139, 111)),
]
for i, (num, name, body, color) in enumerate(phases):
    y = 2.1 + i * 0.82
    pill(slide, num, 0.8, y, 0.55, color)
    text(slide, name, 1.65, y - 0.02, 2.25, 0.28, 15, INK, True)
    text(slide, body, 4.2, y - 0.02, 7.6, 0.28, 14, MUTED)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.08), Inches(y + 0.4), Inches(0.03), Inches(0.45))
    line.fill.solid(); line.fill.fore_color.rgb = RGBColor(210, 223, 224); line.line.fill.background()
footer(slide, 4)

# 5 phase 1
slide = prs.slides.add_slide(prs.slide_layouts[6]); bg(slide); title(slide, "Phase 1", "The working governance workspace", "A practical operating surface for the first line of governance work.")
card(slide, "Regulation Repository", "Track jurisdictions, categories, effective dates, articles, clauses, and source URLs.", 0.75, 2.1, 5.7, 1.35, TEAL)
card(slide, "Compliance Dashboard", "View risk levels, assessment status, policy status, average score, and open findings.", 6.9, 2.1, 5.7, 1.35, ORANGE)
card(slide, "Policy Management", "Create, review, approve, retire, and search governance policy content.", 0.75, 4.0, 5.7, 1.35, RGBColor(91, 105, 190))
card(slide, "Research Workspace", "Ask grounded questions, inspect sources, and save research notes.", 6.9, 4.0, 5.7, 1.35, NAVY)
footer(slide, 5)

# 6 phase 2
slide = prs.slides.add_slide(prs.slide_layouts[6]); bg(slide); title(slide, "Phase 2", "Knowledge becomes searchable evidence", "Every answer can point back to a source chunk.")
steps = [("1", "Ingest", "TXT / Markdown / PDF / DOCX"), ("2", "Chunk", "Recursive overlapping chunks"), ("3", "Embed", "OpenAI embedding model"), ("4", "Retrieve", "Similarity + relevance score"), ("5", "Cite", "Numbered source markers")]
for i, (num, name, body) in enumerate(steps):
    x = 0.65 + i * 2.52
    pill(slide, num, x, 2.25, 0.48, [TEAL, ORANGE, RGBColor(91, 105, 190), NAVY, RGBColor(37, 139, 111)][i])
    text(slide, name, x, 2.85, 2.1, 0.3, 16, INK, True)
    text(slide, body, x, 3.3, 2.1, 0.62, 12, MUTED)
    if i < 4: text(slide, "→", x + 2.12, 2.32, 0.3, 0.3, 18, TEAL, True)
card(slide, "Knowledge domains", "Regulations  •  Industry Standards  •  Audit Findings  •  Governance Documents", 1.0, 5.0, 11.25, 1.0, TEAL)
footer(slide, 6)

# 7 phase 3
slide = prs.slides.add_slide(prs.slide_layouts[6]); bg(slide); title(slide, "Phase 3", "Context-aware agents with real tools", "Specialists use structured database tools and RAG search instead of guessing.")
agent_data = [("Legal Research", "Regulations + policies", TEAL), ("Compliance", "Frameworks + controls", ORANGE), ("Risk Evaluation", "Risk profiles + findings", RGBColor(91, 105, 190)), ("Governance Advisor", "Policies + recommendations", NAVY)]
for i, (name, body, color) in enumerate(agent_data):
    x = 0.7 + (i % 2) * 6.1; y = 2.2 + (i // 2) * 1.9
    card(slide, name, body, x, y, 5.55, 1.25, color)
text(slide, "Shared memory is scoped by session; the orchestrator routes each question to one best-fit specialist.", 0.8, 6.2, 11.5, 0.4, 16, INK, True)
footer(slide, 7)

# 8 phase 4
slide = prs.slides.add_slide(prs.slide_layouts[6]); bg(slide); title(slide, "Phase 4", "MCP makes governance interoperable", "External clients can call the same governed capabilities through Model Context Protocol.")
mcp = [("Policy Repository", "CRUD + semantic policy search"), ("Compliance Tool", "Frameworks, controls, assessments"), ("Audit System", "Audits, findings, remediation"), ("Document Management", "Ingest, search, delete")]
for i, (name, body) in enumerate(mcp):
    x = 0.8 + (i % 2) * 6.1; y = 2.25 + (i // 2) * 1.85
    card(slide, name, body, x, y, 5.5, 1.2, [TEAL, ORANGE, RGBColor(91, 105, 190), NAVY][i])
text(slide, "VS Code / Claude Desktop / compatible agents → stdio MCP servers → shared platform data", 0.8, 6.15, 11.7, 0.35, 16, INK, True)
footer(slide, 8)

# 9 phase 5
slide = prs.slides.add_slide(prs.slide_layouts[6]); bg(slide); title(slide, "Phase 5", "A recommendation with a chain of custody", "The final advice is not a black box: every specialist step is inspectable.")
chain = [("Research", TEAL), ("Compliance", ORANGE), ("Risk", RGBColor(91, 105, 190)), ("Governance", NAVY)]
for i, (name, color) in enumerate(chain):
    x = 0.85 + i * 3.05
    pill(slide, name, x, 2.45, 1.65, color)
    if i < 3: text(slide, "→", x + 2.0, 2.48, 0.6, 0.3, 22, TEAL, True, align=PP_ALIGN.CENTER)
    text(slide, ["Find applicable sources", "Map controls and gaps", "Assess risk and findings", "Prioritize actions"][i], x - 0.05, 3.1, 2.0, 0.55, 12, MUTED, align=PP_ALIGN.CENTER)
card(slide, "Traceability record", "Run ID • agent output • tool/LLM events • durations • final recommendation • persisted audit trail", 1.2, 4.85, 10.9, 1.0, ORANGE)
footer(slide, 9)

# 10 demo flow
slide = prs.slides.add_slide(prs.slide_layouts[6]); bg(slide); title(slide, "Demo flow", "A five-minute story for stakeholders", "Start with a real governance question and follow the evidence to the decision.")
flow = [
    ("Ask", "Should we deploy the credit model in the EU?"),
    ("Research", "Find EU AI Act obligations and standards"),
    ("Assess", "Map controls and current compliance score"),
    ("Evaluate", "Surface high-risk classification and findings"),
    ("Recommend", "Prioritized action plan with trace"),
]
for i, (name, body) in enumerate(flow):
    y = 1.95 + i * 0.82
    pill(slide, str(i + 1), 0.95, y, 0.45, [TEAL, ORANGE, RGBColor(91, 105, 190), NAVY, RGBColor(37, 139, 111)][i])
    text(slide, name, 1.75, y, 1.25, 0.26, 15, INK, True)
    text(slide, body, 3.35, y, 8.8, 0.28, 14, MUTED)
footer(slide, 10)

# 11 operational readiness
slide = prs.slides.add_slide(prs.slide_layouts[6]); bg(slide); title(slide, "Operational readiness", "What is working today", "The POC is deliberately demonstrable, testable, and extensible.")
card(slide, "Validated", "23 backend tests pass. Frontend production build passes across all 16 routes. Auth and protected API flows are covered.", 0.75, 2.15, 5.7, 1.4, TEAL)
card(slide, "Configurable", "OpenAI and LangSmith are environment-driven. Chroma supports a server or embedded local fallback.", 6.9, 2.15, 5.7, 1.4, ORANGE)
card(slide, "Presentable", "Professional workspace shell, domain-focused document ingestion, agent collaboration trace, and MCP configuration.", 0.75, 4.15, 5.7, 1.4, RGBColor(91, 105, 190))
card(slide, "Next hardening", "Add production auth provider, persistent LangGraph checkpointing, migration revisions, and CI/CD deployment controls.", 6.9, 4.15, 5.7, 1.4, NAVY)
footer(slide, 11)

# 12 close
slide = prs.slides.add_slide(prs.slide_layouts[6]); bg(slide, NAVY)
pill(slide, "THE TAKEAWAY", 0.8, 0.85, 1.45, TEAL)
text(slide, "Governance decisions\nwith evidence attached.", 0.8, 1.75, 8.8, 1.35, 40, WHITE, True)
text(slide, "A working foundation for safer, more explainable, and more accountable AI adoption.", 0.85, 3.55, 7.8, 0.6, 20, MINT)
text(slide, "AI Governance & Compliance Platform", 0.85, 6.35, 6.5, 0.3, 14, RGBColor(180, 213, 209), True)
text(slide, "12", 11.9, 6.25, 0.7, 0.5, 18, TEAL, True, align=PP_ALIGN.RIGHT)

prs.save(OUT)
print(OUT)
